"""Real-format, offline badge attribution regression tests (bundled Python)."""
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "scripts/assets/company-badge"
spec = importlib.util.spec_from_file_location("badge_metadata", ASSETS / "badge_metadata.py")
badge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(badge)


class BadgeTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="rufa-badge-formats-")
        self.addCleanup(self.directory.cleanup)
        self.work = Path(self.directory.name)

    def run_helper(self, path, name="alice", sources=None, operation="stamp"):
        request = {"operation": operation, "path": str(path), "username": name, "sources": [str(p) for p in sources or []], "deferPublish": True}
        result = subprocess.run([sys.executable, "-I", "-B", str(ASSETS / "badge_metadata.py")], input=json.dumps(request, ensure_ascii=False) + '\n{"publish":true}\n', text=True, encoding="utf-8", capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        rows = [json.loads(line) for line in result.stdout.splitlines()]
        if operation == "stamp":
            self.assertEqual(rows[0], {"ready": True})
        return rows[-1]

    def office(self, extension, image):
        path = self.work / ("交付文件." + extension)
        if extension == "docx":
            from docx import Document
            document = Document()
            document.add_paragraph("项目报价（保留正文）")
            document.add_picture(str(image))
            document.save(path)
        elif extension == "xlsx":
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image
            workbook = Workbook()
            workbook.active["A1"] = "项目报价"
            workbook.active["A2"] = "=1+2"
            workbook.active.add_image(Image(str(image)), "B3")
            workbook.save(path)
        else:
            from pptx import Presentation
            from pptx.util import Inches
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "项目报价"
            slide.shapes.add_picture(str(image), 0, Inches(2))
            presentation.save(path)
        return path

    def test_png_svg_add_two_users_deduplicate_without_visual_changes(self):
        from PIL import Image
        for extension in ("png", "svg"):
            with self.subTest(extension=extension):
                source = ASSETS / ("rufa-badge." + extension)
                path = self.work / source.name
                shutil.copyfile(source, path)
                original = path.read_bytes()
                self.assertEqual(self.run_helper(path)["authors"], ["alice"])
                self.assertEqual(self.run_helper(path, "张三")["authors"], ["alice", "张三"])
                after = path.read_bytes()
                self.assertFalse(self.run_helper(path, "张三")["changed"])
                self.assertEqual(path.read_bytes(), after)
                self.assertEqual(self.run_helper(path, operation="read")["authors"], ["alice", "张三"])
                if extension == "png":
                    before_image, after_image = Image.open(io.BytesIO(original)), Image.open(path)
                    self.assertEqual(before_image.tobytes(), after_image.tobytes())
                    self.assertEqual([c[2] for c in badge.png_chunks(original)], [c[2] for c in badge.png_chunks(after) if c[0] != b"iTXt"])
                else:
                    self.assertEqual(badge.SVG_TAG.sub(b"", after), original)

    def test_real_office_files_merge_embedded_authors_keep_document_bytes(self):
        image = self.work / "badge.png"
        shutil.copyfile(ASSETS / "rufa-badge.png", image)
        self.run_helper(image, "alice")
        for extension in ("docx", "xlsx", "pptx"):
            with self.subTest(extension=extension):
                path = self.office(extension, image)
                with zipfile.ZipFile(path) as archive:
                    original = {name: archive.read(name) for name in archive.namelist()}
                self.assertEqual(self.run_helper(path, "bob")["authors"], ["alice", "bob"])
                self.assertFalse(self.run_helper(path, "bob")["changed"])
                with zipfile.ZipFile(path) as archive:
                    after = {name: archive.read(name) for name in archive.namelist()}
                for name, body in original.items():
                    if name in ("_rels/.rels", "[Content_Types].xml") or "/media/" in name:
                        continue
                    self.assertEqual(after[name], body, name)
                self.assertIn(b"RufaBadge", after["docProps/custom.xml"])
                for name, body in after.items():
                    if "/media/" in name and name.endswith(".png"):
                        self.assertEqual(badge.png_authors(body), ["alice", "bob"])
                # A round-trip through each actual document library must retain at least one copy.
                if extension == "docx":
                    from docx import Document
                    doc = Document(path)
                    self.assertEqual(doc.paragraphs[0].text, "项目报价（保留正文）")
                    doc.save(path)
                elif extension == "xlsx":
                    from openpyxl import load_workbook
                    doc = load_workbook(path)
                    self.assertEqual(doc.active["A2"].value, "=1+2")
                    doc.save(path)
                else:
                    from pptx import Presentation
                    doc = Presentation(path)
                    self.assertEqual(doc.slides[0].shapes[0].text, "项目报价")
                    doc.save(path)
                self.assertEqual(self.run_helper(path, operation="read")["authors"], ["alice", "bob"])

    def test_pdf_preserves_rendered_pixels_and_existing_document_metadata(self):
        import pymupdf as fitz
        path = self.work / "report.pdf"
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_text((72, 72), "Quote 123.45")
            page.insert_image(fitz.Rect(72, 600, 252, 628), filename=str(ASSETS / "rufa-badge.png"))
            doc.set_metadata({"author": "Original document author", "title": "Quote"})
            doc.save(path)
        with fitz.open(path) as doc:
            pixels = doc[0].get_pixmap().samples
            metadata = doc.metadata
            text = doc[0].get_text()
        self.run_helper(path, "alice")
        self.assertEqual(self.run_helper(path, "bob")["authors"], ["alice", "bob"])
        self.assertFalse(self.run_helper(path, "bob")["changed"])
        with fitz.open(path) as doc:
            self.assertEqual(doc[0].get_pixmap().samples, pixels)
            self.assertEqual(doc[0].get_text(), text)
            self.assertEqual(doc.metadata, metadata)

    def test_sources_survive_format_conversion_hidden_markup_does_not_change_body(self):
        source = self.work / "source.png"
        shutil.copyfile(ASSETS / "rufa-badge.png", source)
        self.run_helper(source, "alice")
        for extension in ("html", "md"):
            path = self.work / ("page." + extension)
            body = '<h1>项目报价</h1><img src="source.png" alt="如发机械有限公司">' if extension == "html" else '# 项目报价\n![如发机械有限公司](source.png)'
            path.write_text(body, encoding="utf-8")
            original = path.read_bytes()
            self.assertEqual(self.run_helper(path, "bob", [source])["authors"], ["alice", "bob"])
            raw = path.read_bytes()
            self.assertNotIn(b"alice", raw)
            self.assertEqual(badge.COMMENT.sub(b"", raw).strip(), original)

    def start_prepared(self, path, sources=None):
        process = subprocess.Popen([sys.executable, "-I", "-B", str(ASSETS / "badge_metadata.py")], stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8")
        self.addCleanup(lambda: process.poll() is None and process.kill())
        process.stdin.write(json.dumps({"operation": "stamp", "path": str(path), "username": "alice", "sources": [str(p) for p in sources or []], "deferPublish": True}) + "\n")
        process.stdin.flush()
        self.assertEqual(json.loads(process.stdout.readline()), {"ready": True})
        return process

    def test_cancel_and_concurrent_changes_do_not_overwrite(self):
        path = self.work / "badge.png"
        shutil.copyfile(ASSETS / "rufa-badge.png", path)
        before = path.read_bytes()
        process = self.start_prepared(path)
        _, _ = process.communicate('{"publish":false}\n', timeout=10)
        self.assertNotEqual(process.returncode, 0)
        self.assertEqual(path.read_bytes(), before)
        source = self.work / "source.md"
        source.write_text("Original", encoding="utf-8")
        process = self.start_prepared(path, [source])
        source.write_text("Updated by another program", encoding="utf-8")
        _, error = process.communicate('{"publish":true}\n', timeout=10)
        self.assertNotEqual(process.returncode, 0)
        self.assertIn("已变化", error)
        self.assertEqual(path.read_bytes(), before)
        process = self.start_prepared(path)
        path.write_bytes(b"concurrent change")
        process.communicate('{"publish":true}\n', timeout=10)
        self.assertNotEqual(process.returncode, 0)
        self.assertEqual(path.read_bytes(), b"concurrent change")
        self.assertEqual(list(self.work.glob(".rufa-badge-*.tmp")), [])

    @unittest.skipUnless(os.name == "nt", "Windows DACL")
    def test_windows_temporary_and_final_file_preserve_restricted_dacl(self):
        path = self.work / "restricted.png"
        shutil.copyfile(ASSETS / "rufa-badge.png", path)
        def powershell(script):
            result = subprocess.run(["pwsh.exe", "-NoProfile", "-NonInteractive", "-Command", "$ErrorActionPreference='Stop'; " + script], capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(result.returncode, 0, result.stderr)
            return result.stdout.strip()
        quote = lambda value: "'" + str(value).replace("'", "''") + "'"
        powershell("$p=" + quote(path) + "; $a=Get-Acl -LiteralPath $p; $a.SetAccessRuleProtection($true,$false); $u=[Security.Principal.WindowsIdentity]::GetCurrent().User; $r=New-Object Security.AccessControl.FileSystemAccessRule($u,'FullControl','Allow'); $a.SetAccessRule($r); Set-Acl -LiteralPath $p -AclObject $a")
        sddl = lambda value: powershell("(Get-Acl -LiteralPath " + quote(value) + ").GetSecurityDescriptorSddlForm([Security.AccessControl.AccessControlSections]::Access)")
        before = sddl(path)
        process = self.start_prepared(path)
        staged = list(self.work.glob(".rufa-badge-*.tmp"))
        self.assertEqual(len(staged), 1)
        staged_sddl = sddl(staged[0])
        self.assertIn("P", staged_sddl.split("(", 1)[0], "staging DACL must be protected")
        self.assertEqual(staged_sddl[staged_sddl.index("("):], before[before.index("("):], "staging cannot inherit broader directory permissions")
        process.communicate('{"publish":true}\n', timeout=10)
        self.assertEqual(process.returncode, 0)
        self.assertEqual(sddl(path), before)

    def test_corrupt_metadata_unsupported_files_and_signed_documents_are_preserved(self):
        path = self.work / "corrupt.md"
        path.write_text("private content\n<!-- rufa-badge:broken -->", encoding="utf-8")
        original = path.read_bytes()
        request = json.dumps({"operation": "stamp", "path": str(path), "username": "bob", "deferPublish": True})
        result = subprocess.run([sys.executable, "-I", "-B", str(ASSETS / "badge_metadata.py")], input=request + '\n{"publish":true}\n', capture_output=True, text=True, encoding="utf-8")
        self.assertNotEqual(result.returncode, 0)
        self.assertNotIn("private content", result.stdout + result.stderr)
        self.assertEqual(path.read_bytes(), original)
        with self.assertRaises(badge.BadgeError):
            badge.read_authors(b"anything", ".exe")
        doc = self.office("docx", ASSETS / "rufa-badge.png")
        with zipfile.ZipFile(doc, "a") as archive:
            archive.writestr("_xmlsignatures/sig1.xml", "<signature/>")
        with self.assertRaises(badge.BadgeError):
            badge.read_authors(doc.read_bytes(), ".docx")


if __name__ == "__main__":
    unittest.main(verbosity=2)
