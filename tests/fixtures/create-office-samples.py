"""Synthetic plugin smoke fixtures; no employee documents are read."""
from pathlib import Path
import sys
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
import pymupdf

root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
text = "Rufa office test 2026 - local document extraction"
document = Document()
document.add_paragraph("如发机械插件验证 " + text)
document.save(root / "sample.docx")
workbook = Workbook()
workbook.active.append(["如发机械", text])
workbook.active["A2"] = 10
workbook.active["A3"] = 20
workbook.active["A4"] = "=SUM(A2:A3)"
workbook.save(root / "sample.xlsx")
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
slide.shapes.add_textbox(500000, 500000, 8000000, 1500000).text = "如发机械 " + text
presentation.save(root / "sample.pptx")
pdf = pymupdf.open()
page = pdf.new_page(width=800, height=220)
page.insert_text((35, 80), text, fontsize=25)
page.insert_text((35, 130), "如发机械本地文字识别", fontname="china-s", fontsize=25)
pdf.save(root / "sample.pdf")
image = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
image.save(root / "scan.png")
scan = pymupdf.open()
scan.new_page(width=800, height=220).insert_image(pymupdf.Rect(0, 0, 800, 220), stream=image.tobytes("png"))
scan.save(root / "scanned.pdf")
scan.close()
pdf.close()
(root / "sample.txt").write_text("如发机械 " + text, encoding="utf-8")
print("Synthetic Office/PDF/OCR fixtures ready")
