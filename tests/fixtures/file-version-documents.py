import os
import sys
import zipfile
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

directory = sys.argv[1]
os.makedirs(directory, exist_ok=True)
for index in (1, 2):
    doc = Document()
    doc.add_paragraph('项目报价说明')
    doc.add_paragraph('交货日期：' + ('11月1日' if index == 1 else '11月8日'))
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '合同金额'
    table.cell(0, 1).text = '1000' if index == 1 else '1200'
    doc.save(os.path.join(directory, f'word-{index}.blob'))
    book = Workbook()
    sheet = book.active
    sheet.title = '销售'
    sheet['A1'] = '金额'
    sheet['B2'] = 1000 if index == 1 else 1200
    sheet['C2'] = '=B2*1.1' if index == 1 else '=B2*1.2'
    sheet['D2'] = '保持原样'
    book.save(os.path.join(directory, f'excel-{index}.blob'))
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = '项目汇报'
    slide.placeholders[1].text = '本周完成 3 项' if index == 1 else '本周完成 5 项'
    deck.save(os.path.join(directory, f'slides-{index}.blob'))

# A separately saved spreadsheet represents a colleague's newer manual edit.
book.active['B2'] = 1500
book.active['C2'] = '=B2*1.3'
book.save(os.path.join(directory, 'excel-external.blob'))

with zipfile.ZipFile(os.path.join(directory, 'untrusted.blob'), 'w') as file:
    file.writestr('word/document.xml', '<!DOCTYPE a [<!ENTITY x SYSTEM "file:///private">]><a>&x;</a>')

# Bounded extraction rejects an oversized declaration before inflating it.
with zipfile.ZipFile(os.path.join(directory, 'oversized.blob'), 'w', compression=zipfile.ZIP_DEFLATED) as file:
    file.writestr('word/document.xml', b'x' * (33 * 1024 * 1024))
